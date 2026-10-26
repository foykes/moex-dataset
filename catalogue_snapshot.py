"""A4-01: unqualified catalogue candidates, without current-target I/O.

Profiles are local consumer policies, not verified MOEX source contracts.
openpyxl uses temporary worksheet files in the caller-owned TMP.
"""
import csv
import datetime
import errno
import hashlib
import io
import re
import sys
import time
from itertools import zip_longest
from types import SimpleNamespace

BODY_LIMIT = 16 * 1024**2
COLUMN_LIMIT = 256
ROW_LIMIT = 100000
PARSED_CELL_LIMIT = 1000000
PAYLOAD_LIMIT = 64 * 1024**2
CELL_UTF16_LIMIT = 32767

_PROFILES = ('A4_CATALOGUE_EXPLICIT_ENCODING_V1', 'A4_CATALOGUE_LEGACY_CP1251_V1')
_ARTIFACTS = ('moex_full.csv', 'moex_full.xlsx', 'moex_stocks.csv', 'moex_stocks.xlsx')
_ID = re.compile(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,127}', re.ASCII)
_UTC = re.compile(r'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?(?:Z|\+00:00)', re.ASCII)
_MIME = re.compile(r'text/csv(?:[ \t]*;[ \t]*charset[ \t]*=[ \t]*(?:"([a-z0-9-]+)"|([a-z0-9-]+)))?', re.ASCII | re.I)
_REASONS = {
    'A4_INPUT', 'A4_HTTP_STATUS', 'A4_CONTENT_TYPE', 'A4_CHARSET', 'A4_BOM',
    'A4_SCHEMA', 'A4_EMPTY_UNIVERSE', 'A4_DECODE', 'A4_CSV_PARSE',
    'A4_BUDGET_BODY_BYTES', 'A4_BUDGET_COLUMNS', 'A4_BUDGET_ROWS',
    'A4_BUDGET_PARSED_CELLS', 'A4_BUDGET_PAYLOAD_BYTES', 'A4_XLSX_TEXT_LIMIT',
    'A4_XLSX_VALUE', 'A4_READBACK', 'A4_MEMORY', 'A4_DISK_FULL', 'A4_IO',
    'A4_RENDER', 'A4_PREPARE', 'A4_LOGGING_INCOMPLETE',
}
_XML_INVALID = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def _refuse(reason, artifact):
    error = ValueError(reason)
    error.a4_failure = dict(version=1, reason=reason, artifact=artifact)
    raise error


def _failure(error, reason, artifact):
    # Не выводим args/str exporter exception: там может быть вся исходная ячейка.
    try:
        previous = getattr(error, 'a4_failure', None)
        if (type(error) is ValueError and type(previous) is dict and set(previous) == {'version', 'reason', 'artifact'}
                and type(previous.get('version')) is int and previous['version'] == 1
                and type(previous.get('reason')) is str and previous['reason'] in _REASONS
                and previous.get('artifact') in (*_ARTIFACTS, 'SOURCE_BODY', 'CATALOGUE_FRAME')):
            return dict(previous)
        if isinstance(error, MemoryError):
            reason = 'A4_MEMORY'
        elif isinstance(error, OSError):
            reason = 'A4_DISK_FULL' if error.errno == errno.ENOSPC else 'A4_IO'
        elif isinstance(error, UnicodeDecodeError):
            reason = 'A4_DECODE'
        elif isinstance(error, csv.Error):
            reason = 'A4_CSV_PARSE'
        elif type(error) is RuntimeError and error.args == ('A4_LOGGING_INCOMPLETE',):
            reason = 'A4_LOGGING_INCOMPLETE'
        result = dict(version=1, reason=reason, artifact=artifact)
        error.a4_failure = dict(result)
        return result
    except BaseException:
        # Ошибка добавления evidence никогда не заменяет primary.
        return dict(version=1, reason=reason, artifact=artifact)


def _identity(value):
    return type(value) is str and _ID.fullmatch(value) is not None


def _inputs(content, profile, source_ref, snapshot_id, captured_at, artifact_ids):
    if (type(content) is not bytes or type(profile) is not str or profile not in _PROFILES
            or not _identity(source_ref) or not _identity(snapshot_id)
            or type(captured_at) is not str or _UTC.fullmatch(captured_at) is None):
        _refuse('A4_INPUT', 'SOURCE_BODY')
    try:
        timestamp = datetime.datetime.fromisoformat(captured_at.replace('Z', '+00:00'))
    except ValueError:
        _refuse('A4_INPUT', 'SOURCE_BODY')
    if timestamp.utcoffset() != datetime.timedelta(0):
        _refuse('A4_INPUT', 'SOURCE_BODY')
    if artifact_ids is not None and (type(artifact_ids) is not dict
            or any(name not in _ARTIFACTS or not _identity(value)
                   for name, value in artifact_ids.items())):
        _refuse('A4_INPUT', 'SOURCE_BODY')
    return timestamp.isoformat().replace('+00:00', 'Z')


def _codec(status, content_type, profile):
    if type(status) is not int or status != 200:
        _refuse('A4_HTTP_STATUS', 'SOURCE_BODY')
    if type(content_type) is not str or len(content_type) > 256:
        _refuse('A4_CONTENT_TYPE', 'SOURCE_BODY')
    match = _MIME.fullmatch(content_type.strip(' \t'))
    if match is None:
        _refuse('A4_CONTENT_TYPE', 'SOURCE_BODY')
    charset = match.group(1) or match.group(2)
    if charset is None:
        if profile != 'A4_CATALOGUE_LEGACY_CP1251_V1':
            _refuse('A4_CHARSET', 'SOURCE_BODY')
        return 'cp1251'
    codecs = {'cp1251': 'cp1251', 'windows-1251': 'cp1251', 'utf-8': 'utf-8', 'utf8': 'utf-8'}
    codec = codecs.get(charset.lower())
    if codec is None:
        _refuse('A4_CHARSET', 'SOURCE_BODY')
    return codec


def _text_cell(value):
    # openpyxl иначе молча обрезает строку и автоматически выводит f/e типы.
    if _XML_INVALID.search(value) is not None:
        _refuse('A4_XLSX_VALUE', 'CATALOGUE_FRAME')
    if len(value.encode('utf-16-le')) // 2 > CELL_UTF16_LIMIT:
        _refuse('A4_XLSX_TEXT_LIMIT', 'CATALOGUE_FRAME')


def _parse(content, codec):
    if len(content) > BODY_LIMIT:
        _refuse('A4_BUDGET_BODY_BYTES', 'SOURCE_BODY')
    if content.startswith(b'\xef\xbb\xbf'):
        if codec != 'utf-8':
            _refuse('A4_BOM', 'SOURCE_BODY')
        content = content[3:]
        if content.startswith(b'\xef\xbb\xbf'):
            _refuse('A4_BOM', 'SOURCE_BODY')
    elif content.startswith((b'\xff\xfe', b'\xfe\xff', b'\x00\x00\xfe\xff')):
        _refuse('A4_BOM', 'SOURCE_BODY')
    reader = csv.reader(io.StringIO(content.decode(codec, errors='strict'), newline=''), strict=True)
    columns = next(reader, None)
    if columns is None:
        _refuse('A4_EMPTY_UNIVERSE', 'SOURCE_BODY')
    if len(columns) > COLUMN_LIMIT:
        _refuse('A4_BUDGET_COLUMNS', 'CATALOGUE_FRAME')
    if len(columns) > PARSED_CELL_LIMIT:
        _refuse('A4_BUDGET_PARSED_CELLS', 'CATALOGUE_FRAME')
    if (not columns or any(not value.strip() for value in columns)
            or len(set(columns)) != len(columns)
            or any(columns.count(name) != 1 for name in ('TRADE_CODE', 'SUPERTYPE'))):
        _refuse('A4_SCHEMA', 'CATALOGUE_FRAME')
    for value in columns:
        _text_cell(value)
    rows, cells, usable = [], len(columns), False
    code_position = columns.index('TRADE_CODE')
    for row in reader:
        if len(row) != len(columns):
            _refuse('A4_SCHEMA', 'CATALOGUE_FRAME')
        if len(rows) >= ROW_LIMIT:
            _refuse('A4_BUDGET_ROWS', 'CATALOGUE_FRAME')
        if cells + len(row) > PARSED_CELL_LIMIT:
            _refuse('A4_BUDGET_PARSED_CELLS', 'CATALOGUE_FRAME')
        for value in row:
            _text_cell(value)
        usable = usable or bool(row[code_position].strip())
        rows.append(row)
        cells += len(row)
    if not usable:
        _refuse('A4_EMPTY_UNIVERSE', 'CATALOGUE_FRAME')
    # A1 processing admission остаётся обязанностью будущего caller.
    import pandas as pd
    full = pd.DataFrame(rows, columns=columns, index=range(1, len(rows) + 1))
    stocks = full[full['SUPERTYPE'].isin(('Акции', 'Депозитарные расписки'))].copy()
    stocks.reset_index(drop=True, inplace=True)
    return full, stocks


def _payload_buffer(artifact, budget):
    """Charge aggregate high-water before extending a seekable payload buffer."""
    stream = io.BytesIO()
    original_write, original_truncate = stream.write, stream.truncate
    high_water = 0

    def reserve(end):
        nonlocal high_water
        growth = max(0, end - high_water)
        if budget['total'] + growth > PAYLOAD_LIMIT:
            _refuse('A4_BUDGET_PAYLOAD_BYTES', artifact)
        budget['total'] += growth
        high_water += growth

    def write(value):
        size = memoryview(value).nbytes
        reserve(stream.tell() + size if size else high_water)
        return original_write(value)

    def truncate(size=None):
        reserve(stream.tell() if size is None else size)
        return original_truncate(size)

    stream.write, stream.truncate = write, truncate
    return stream


def _render_csv(frame, artifact, budget):
    with _payload_buffer(artifact, budget) as stream:
        def write_text(value):
            return stream.write(value.encode('utf-8'))
        writer = csv.writer(SimpleNamespace(write=write_text), lineterminator='\r\n')
        writer.writerow([''] + list(frame.columns))
        for row in frame.itertuples(index=True, name=None):
            writer.writerow(row)
        return stream.getvalue()


def _close_book(book, primary):
    try:
        book.close()
    except BaseException:
        if not primary:
            raise


def _render_xlsx(frame, artifact, budget):
    import openpyxl
    book = openpyxl.Workbook()
    failed = True
    try:
        sheet = book.active
        sheet.title = 'Sheet1'
        for column, value in enumerate(frame.columns, 2):
            cell = sheet.cell(1, column, value)
            cell.data_type = 's'
        for row_number, row in enumerate(frame.itertuples(index=True, name=None), 2):
            sheet.cell(row_number, 1, row[0])
            for column, value in enumerate(row[1:], 2):
                cell = sheet.cell(row_number, column, value)
                cell.data_type = 's'
        with _payload_buffer(artifact, budget) as stream:
            book.save(stream)
            result = stream.getvalue()
        failed = False
        return result
    finally:
        _close_book(book, failed)


def _validate_csv(payload, frame, artifact):
    actual = csv.reader(io.StringIO(payload.decode('utf-8'), newline=''), strict=True)
    if next(actual, None) != [''] + list(frame.columns):
        _refuse('A4_READBACK', artifact)
    sentinel = object()
    for received, expected in zip_longest(actual, frame.itertuples(index=True, name=None), fillvalue=sentinel):
        if received is sentinel or expected is sentinel or received != [str(expected[0]), *expected[1:]]:
            _refuse('A4_READBACK', artifact)


def _validate_xlsx(payload, frame, artifact):
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(payload), data_only=False)
    failed = True
    try:
        if book.sheetnames != ['Sheet1']:
            _refuse('A4_READBACK', artifact)
        sheet = book['Sheet1']
        if (sheet.max_row != len(frame) + 1 or sheet.max_column != len(frame.columns) + 1
                or sheet['A1'].value is not None or sheet['A1'].data_type in {'f', 'e'}):
            _refuse('A4_READBACK', artifact)
        for column, expected in enumerate(frame.columns, 2):
            cell = sheet.cell(1, column)
            if cell.value != expected or cell.data_type != 's':
                _refuse('A4_READBACK', artifact)
        for row_number, expected in enumerate(frame.itertuples(index=True, name=None), 2):
            cell = sheet.cell(row_number, 1)
            if type(cell.value) is not int or cell.value != expected[0] or cell.data_type != 'n':
                _refuse('A4_READBACK', artifact)
            for column, value in enumerate(expected[1:], 2):
                cell = sheet.cell(row_number, column)
                if (cell.data_type in {'f', 'e'} or ('' if cell.value is None else cell.value) != value
                        or value and (type(cell.value) is not str or cell.data_type != 's')):
                    _refuse('A4_READBACK', artifact)
        failed = False
    finally:
        _close_book(book, failed)


def _logging_fault():
    raise RuntimeError('A4_LOGGING_INCOMPLETE') from None


def _healthy(info):
    if info['context'] is None:
        return
    try:
        import run_logging
        context = info['context']
        if type(context) is not dict or context.get('worker_id', 'worker') is not None:
            _logging_fault()
        levels = {'DEBUG': 10, 'INFO': 20, 'WARNING': 30, 'ERROR': 40}
        config = context.get('config')
        if (type(config) is not dict or levels.get(config.get('file_level'), 100) > 20
                or levels.get(config.get('console_level'), 100) > 20):
            _logging_fault()
        health = run_logging.check_logging_health(context)
        if type(health) is not dict or health.get('healthy') is not True or health.get('state') != 'ACTIVE':
            _logging_fault()
    except BaseException:
        _logging_fault()


def _event(info, name, *, error=None):
    if info['context'] is None:
        return
    try:
        import run_logging
        _healthy(info)
        level = 'ERROR' if error is not None else 'INFO'
        fields = dict(stage='data_gathering', fields={}, snapshot_id=info.get('snapshot_id'),
            elapsed_ms=max(0, (time.perf_counter() - info['started']) * 1000),
            outcome='FAILED' if error is not None else 'PREPARED')
        artifact = info['artifact']
        if artifact in _ARTIFACTS:
            fields.update(file='ticker_lists/' + artifact,
                          artifact_id=info['artifact_ids'].get(artifact))
        if info.get('rows') is not None:
            fields['counts'] = dict(rows_received=info['rows'], rows_accepted=info['rows'],
                rows_quarantined=None, pages=None, instruments=None,
                null_reasons={'rows_quarantined': 'NOT_ASSESSED', 'pages': 'NOT_APPLICABLE',
                              'instruments': 'NOT_ASSESSED'})
        schema = ('REJECTED' if error is not None and error['reason'] in
                  {'A4_SCHEMA', 'A4_EMPTY_UNIVERSE'} else info.get('schema', 'UNKNOWN'))
        message = 'Каталог MOEX_LISTING_CSV: status={}; encoding={}; schema={}; {}'.format(
            info.get('status', 'UNKNOWN'), info.get('codec', 'UNKNOWN'), schema, name)
        if error is not None:
            projection = RuntimeError(error['reason'])
            record = run_logging.make_error(info['context'], projection, 'data_gathering', code=error['reason'])
            record['category'] = ('CAPACITY' if error['reason'].startswith('A4_BUDGET_')
                or error['reason'] == 'A4_XLSX_TEXT_LIMIT' else
                'MEMORY_ERROR' if error['reason'] == 'A4_MEMORY' else
                'DISK_FULL' if error['reason'] == 'A4_DISK_FULL' else
                'IO' if error['reason'] in {'A4_IO', 'A4_RENDER'} else
                'QUALITY' if error['reason'] == 'A4_READBACK' else 'SOURCE')
            for key in ('snapshot_id', 'file', 'artifact_id'):
                if fields.get(key) is not None:
                    record[key] = fields[key]
                    record['null_reasons'].pop(key, None)
            fields['error'] = record
        emitted = run_logging.emit_event(info['context'], level, 'a4_' + name, message, fields)
        if type(emitted) is not dict:
            _logging_fault()
        producer, sequence = emitted.get('producer'), emitted.get('sequence')
        if (emitted.get('accepted') is not True or type(producer) is not int
                or not 0 <= producer < 2**64 or type(sequence) is not int
                or sequence <= info.get('last_sequence', 0)
                or info.get('producer', producer) != producer
                or level == 'ERROR' and emitted.get('confirmed') is not True):
            _logging_fault()
        info['producer'], info['last_sequence'] = producer, sequence
        info['sequences'].append(sequence)
        _healthy(info)
    except BaseException:
        _logging_fault()


def _barrier(info):
    if info['context'] is None:
        return dict(status='NOT_CONNECTED', producer=None, barrier_sequence=None, event_sequences=[])
    try:
        import run_logging
        barrier = run_logging.flush_logging(info['context'])
        if (type(barrier) is not dict or barrier.get('accepted') is not True
                or barrier.get('confirmed') is not True
                or type(barrier.get('producer')) is not int
                or barrier.get('producer') != info.get('producer')
                or type(barrier.get('sequence')) is not int
                or barrier['sequence'] <= info.get('last_sequence', 0)):
            _logging_fault()
        _healthy(info)
        return dict(status='CONFIRMED', producer=info['producer'],
                    barrier_sequence=barrier['sequence'], event_sequences=list(info['sequences']))
    except BaseException:
        _logging_fault()


def prepare_catalogue_snapshot(content, *, consumer_profile, status, content_type,
        source_ref, snapshot_id, captured_at_utc, logging_context=None, artifact_ids=None):
    """Prepare four payloads; no source request, current path or activation."""
    info = dict(context=logging_context, artifact='SOURCE_BODY', artifact_ids={},
                started=time.perf_counter(), sequences=[])
    reason = 'A4_PREPARE'
    try:
        captured = _inputs(content, consumer_profile, source_ref, snapshot_id,
                           captured_at_utc, artifact_ids)
        info.update(snapshot_id=snapshot_id, artifact_ids=dict(artifact_ids or {}))
        _healthy(info)
        _event(info, 'started')
        info['status'] = status if type(status) is int and 100 <= status <= 599 else 'UNKNOWN'
        codec = _codec(status, content_type, consumer_profile)
        info.update(status=status, codec=codec)
        _event(info, 'response_accepted')
        full, stocks = _parse(content, codec)
        info.update(artifact='CATALOGUE_FRAME', rows=len(full), schema='ACCEPTED')
        _event(info, 'parsed')
        artifacts, outcomes, budget = {}, {}, {'total': 0}
        for prefix, frame in (('moex_full', full), ('moex_stocks', stocks)):
            for extension, render, validate in (('csv', _render_csv, _validate_csv),
                                                ('xlsx', _render_xlsx, _validate_xlsx)):
                artifact = prefix + '.' + extension
                info.update(artifact=artifact, rows=len(frame))
                reason = 'A4_RENDER'
                payload = render(frame, artifact, budget)
                if type(payload) is not bytes:
                    _refuse('A4_READBACK', artifact)
                reason = 'A4_READBACK'
                validate(payload, frame, artifact)
                artifacts[artifact] = payload
                outcomes[artifact] = dict(bytes=len(payload), sha256=hashlib.sha256(payload).hexdigest(),
                                         render='PASS', readback='PASS')
                _event(info, 'rendered')
        info.update(artifact='CATALOGUE_FRAME', rows=len(full))
        reason = 'A4_PREPARE'
        receipt = dict(receipt_version='A4_CANDIDATE_RECEIPT_V1', consumer_profile=consumer_profile,
            source_ref=source_ref, snapshot_id=snapshot_id, captured_at_utc=captured,
            structural_validation=dict(parsing='PASS', rendering='PASS', readback='PASS'),
            source_contract_qualification={key: 'NOT_VERIFIED' for key in
                ('status', 'encoding', 'full_schema', 'origin_redirects', 'target_live')},
            source_feasibility='NOT_VERIFIED', activation_authority='NOT_GRANTED', encoding=codec,
            ordered_schema=list(full.columns), source_sha256=hashlib.sha256(content).hexdigest(),
            counts=dict(full_rows=len(full), stocks_rows=len(stocks)), artifacts=outcomes)
        info.update(artifact='CATALOGUE_FRAME', rows=len(full))
        _event(info, 'candidate_ready')
        receipt['logging'] = _barrier(info)
        return dict(artifacts=artifacts, receipt=receipt)
    except BaseException as error:
        try:
            observation = _failure(error, reason, info['artifact'])
        except BaseException:
            observation = None
        if logging_context is not None:
            try:
                if observation is None:
                    _logging_fault()
                _event(info, 'failed', error=observation)
            except BaseException:
                marker = 'A4_LOGGING_FAILURE: LOGGING_INCOMPLETE'
                try:
                    print(marker, file=sys.stderr)
                except BaseException:
                    pass
                try:
                    BaseException.add_note(error, marker)
                except BaseException:
                    pass
        raise
