"""A4 candidate semantics, independent format oracles and honest legacy gaps.

MDS_A4_LEGACY_RED=1 changes only the two explicitly named legacy assertions to
the required A4 behavior. Their separate expected failing runs are evidence of
the unchanged caller, not helper failures or permission to contact a source.
"""

import csv
import errno
import hashlib
import io
import json
import os
from pathlib import Path
from types import SimpleNamespace

import pytest


ROOT = Path(__file__).resolve().parents[2]
EXPLICIT = 'A4_CATALOGUE_EXPLICIT_ENCODING_V1'
LEGACY = 'A4_CATALOGUE_LEGACY_CP1251_V1'
NAMES = ('moex_full.csv', 'moex_full.xlsx', 'moex_stocks.csv', 'moex_stocks.xlsx')
WHEN = '2026-10-26T00:00:00Z'


@pytest.fixture
def cases():
    return json.loads((ROOT / 'tests/fixtures/a4/catalogue_cases.json').read_text(encoding='utf-8'))


@pytest.fixture
def candidate(monkeypatch):
    monkeypatch.syspath_prepend(str(ROOT))
    import catalogue_snapshot
    return catalogue_snapshot


def encoded(columns, rows, codec='cp1251'):
    # Fixture construction is deliberately independent of the product renderer.
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, lineterminator='\r\n')
    writer.writerow(columns)
    writer.writerows(rows)
    return stream.getvalue().encode(codec)


def prepare(candidate, content, **changes):
    options = dict(consumer_profile=EXPLICIT, status=200,
                   content_type='text/csv; charset=windows-1251',
                   source_ref='MOEX_LISTING_FIXTURE', snapshot_id='a4-fixture-20261026',
                   captured_at_utc=WHEN)
    options.update(changes)
    return candidate.prepare_catalogue_snapshot(content, **options)


def matrix(columns, rows, start):
    return [[''] + columns] + [[str(index)] + row for index, row in enumerate(rows, start)]


def csv_oracle(payload, columns, rows, start):
    assert type(payload) is bytes
    assert not payload.startswith(b'\xef\xbb\xbf')
    actual = list(csv.reader(io.StringIO(payload.decode('utf-8'), newline=''), strict=True))
    assert actual == matrix(columns, rows, start)


def xlsx_oracle(payload, columns, rows, start):
    # data_only=True would conceal formulas; error cells can retain the same
    # visible value as literal '#N/A', so the type assertions are independent.
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(payload), data_only=False)
    try:
        assert book.sheetnames == ['Sheet1']
        sheet = book['Sheet1']
        assert sheet.max_row == len(rows) + 1
        assert sheet.max_column == len(columns) + 1
        assert sheet['A1'].value is None
        for column, expected in enumerate(columns, 2):
            cell = sheet.cell(1, column)
            assert cell.value == expected and cell.data_type == 's'
        for row_number, (index, expected_row) in enumerate(enumerate(rows, start), 2):
            index_cell = sheet.cell(row_number, 1)
            assert type(index_cell.value) is int and index_cell.value == index
            assert index_cell.data_type == 'n'
            for column, expected in enumerate(expected_row, 2):
                cell = sheet.cell(row_number, column)
                assert ('' if cell.value is None else cell.value) == expected
                assert cell.data_type not in {'f', 'e'}
                if expected:
                    assert type(cell.value) is str and cell.data_type == 's'
    finally:
        book.close()


def candidate_oracle(result, case):
    assert set(result) == {'artifacts', 'receipt'}
    assert set(result['artifacts']) == set(NAMES)
    for prefix, rows, start in [('moex_full', case['rows'], 1),
                                ('moex_stocks', case['stocks'], 0)]:
        csv_oracle(result['artifacts'][prefix + '.csv'], case['columns'], rows, start)
        xlsx_oracle(result['artifacts'][prefix + '.xlsx'], case['columns'], rows, start)


def assert_failure(error, reason=None, artifact=None):
    failure = error.a4_failure
    assert set(failure) == {'version', 'reason', 'artifact'}
    assert failure['version'] == 1
    assert isinstance(failure['reason'], str) and failure['reason'].startswith('A4_')
    assert failure['artifact'] in set(NAMES) | {'SOURCE_BODY', 'CATALOGUE_FRAME'}
    if reason is not None:
        assert failure['reason'] == reason
    if artifact is not None:
        assert failure['artifact'] == artifact


def test_literal_four_payloads_and_independent_receipt(candidate, cases):
    case = cases['mixed']
    content = encoded(case['columns'], case['rows'])
    result = prepare(candidate, content)
    candidate_oracle(result, case)
    receipt = result['receipt']
    assert receipt['receipt_version'] == 'A4_CANDIDATE_RECEIPT_V1'
    assert receipt['consumer_profile'] == EXPLICIT
    assert receipt['snapshot_id'] == 'a4-fixture-20261026'
    assert receipt['captured_at_utc'] == WHEN
    assert receipt['source_ref'] == 'MOEX_LISTING_FIXTURE'
    assert receipt['encoding'] == 'cp1251'
    assert receipt['ordered_schema'] == case['columns']
    assert receipt['source_sha256'] == hashlib.sha256(content).hexdigest()
    assert receipt['counts'] == {'full_rows': 11, 'stocks_rows': 7}
    assert receipt['structural_validation'] == dict(parsing='PASS', rendering='PASS', readback='PASS')
    assert receipt['source_contract_qualification'] == {
        key: 'NOT_VERIFIED' for key in ('status', 'encoding', 'full_schema', 'origin_redirects', 'target_live')}
    assert receipt['source_feasibility'] == 'NOT_VERIFIED'
    assert receipt['activation_authority'] == 'NOT_GRANTED'
    assert set(receipt['artifacts']) == set(NAMES)
    for name, payload in result['artifacts'].items():
        assert receipt['artifacts'][name] == dict(bytes=len(payload), sha256=hashlib.sha256(payload).hexdigest(),
                                                render='PASS', readback='PASS')
    assert receipt['logging']['status'] == 'NOT_CONNECTED'


def test_legacy_undetectable_utf8_remains_unqualified(candidate, cases):
    case = cases['undeclared_utf8']
    result = prepare(candidate, case['source_text'].encode('utf-8'),
                     consumer_profile=LEGACY, content_type='text/csv')
    candidate_oracle(result, case)
    assert result['receipt']['counts'] == {'full_rows': 1, 'stocks_rows': 0}
    assert result['receipt']['encoding'] == 'cp1251'
    assert set(result['receipt']['source_contract_qualification'].values()) == {'NOT_VERIFIED'}


def test_formula_and_error_headers_and_exact_duplicate_rows_are_literal(candidate):
    columns = ['TRADE_CODE', 'SUPERTYPE', '=1+1', '#N/A']
    row = ['SBER', 'Акции', '-1', '@SUM(A1)']
    expected = dict(columns=columns, rows=[list(row), list(row)], stocks=[list(row), list(row)])
    result = prepare(candidate, encoded(columns, [row, row]))
    candidate_oracle(result, expected)
    assert result['receipt']['counts'] == {'full_rows': 2, 'stocks_rows': 2}


@pytest.mark.parametrize('charset', ['cp1251', 'windows-1251', 'CP1251'])
def test_explicit_cp1251_aliases(candidate, cases, charset):
    result = prepare(candidate, cases['simple_text'].encode('cp1251'),
                     content_type='text/csv; charset=' + charset)
    assert result['receipt']['encoding'] == 'cp1251'


@pytest.mark.parametrize('bom', [False, True])
def test_declared_utf8_and_optional_bom(candidate, cases, bom):
    content = cases['simple_text'].encode('utf-8')
    if bom:
        content = b'\xef\xbb\xbf' + content
    result = prepare(candidate, content, content_type='text/csv; charset=utf-8')
    expected = dict(columns=['TRADE_CODE', 'SUPERTYPE'], rows=[['SBER', 'Акции']], stocks=[['SBER', 'Акции']])
    candidate_oracle(result, expected)
    assert result['receipt']['encoding'] == 'utf-8'


@pytest.mark.parametrize('status', [0, 204, 301, 403, 429, 500, True])
def test_status_rejected_before_render(candidate, cases, monkeypatch, status):
    calls = []
    monkeypatch.setattr(candidate, '_render_csv', lambda *args: calls.append(args))
    with pytest.raises(ValueError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), status=status)
    assert_failure(captured.value)
    assert calls == []


@pytest.mark.parametrize('content_type', ['text/html', 'application/json', '', None,
    'text/csv; charset=utf-16', 'text/csv; charset=cp1251; charset=utf-8',
    'text/csv; header=present'])
def test_mime_and_charset_refusal(candidate, cases, content_type):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), content_type=content_type)
    assert_failure(captured.value)


def test_profile_is_explicit_and_profiles_do_not_imply_source_contract(candidate, cases):
    content = cases['simple_text'].encode('cp1251')
    for changes in ({'consumer_profile': 'UNKNOWN'}, {'content_type': 'text/csv'}):
        with pytest.raises(ValueError) as captured:
            prepare(candidate, content, **changes)
        assert_failure(captured.value)
    with pytest.raises(TypeError):
        candidate.prepare_catalogue_snapshot(content, status=200, content_type='text/csv',
            source_ref='MOEX_LISTING_FIXTURE', snapshot_id='fixture', captured_at_utc=WHEN)


@pytest.mark.parametrize('content,content_type', [
    (b'TRADE_CODE,SUPERTYPE\r\nSBER,\x98\r\n', 'text/csv; charset=cp1251'),
    (b'TRADE_CODE,SUPERTYPE\r\nSBER,\xff\r\n', 'text/csv; charset=utf-8')])
def test_strict_decoder_exception_survives(candidate, content, content_type):
    with pytest.raises(UnicodeDecodeError) as captured:
        prepare(candidate, content, content_type=content_type)
    assert_failure(captured.value, artifact='SOURCE_BODY')


def test_bom_cannot_contradict_selected_encoding(candidate, cases):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, b'\xef\xbb\xbf' + cases['simple_text'].encode('utf-8'))
    assert_failure(captured.value, 'A4_BOM', 'SOURCE_BODY')


@pytest.mark.parametrize('columns,rows', [
    (['TRADE_CODE'], [['SBER']]),
    (['TRADE_CODE', 'SUPERTYPE', ''], [['SBER', 'Акции', 'x']]),
    (['TRADE_CODE', 'SUPERTYPE', '  '], [['SBER', 'Акции', 'x']]),
    (['TRADE_CODE', 'SUPERTYPE', 'NAME', 'NAME'], [['SBER', 'Акции', 'a', 'b']]),
    (['TRADE_CODE', 'TRADE_CODE', 'SUPERTYPE'], [['SBER', 'SBER', 'Акции']]),
    (['TRADE_CODE', 'SUPERTYPE'], [['SBER']]),
    (['TRADE_CODE', 'SUPERTYPE'], [['SBER', 'Акции', 'extra']])])
def test_unknown_structure_fails_without_schema_repair(candidate, columns, rows):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, encoded(columns, rows))
    assert_failure(captured.value, 'A4_SCHEMA')


@pytest.mark.parametrize('content', [b'', b'TRADE_CODE,SUPERTYPE\r\n',
    b'TRADE_CODE,SUPERTYPE\r\n,shares\r\n  ,shares\r\n'])
def test_empty_usable_universe_fails(candidate, content):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, content)
    assert_failure(captured.value)


def test_malformed_csv_retains_native_error(candidate, cases):
    for source in cases['invalid_csv']:
        with pytest.raises(csv.Error) as captured:
            prepare(candidate, source.encode('cp1251'))
        assert_failure(captured.value, artifact='SOURCE_BODY')


@pytest.mark.parametrize('field,value', [
    ('snapshot_id', ''), ('snapshot_id', 'bad\nidentifier'),
    ('source_ref', ''), ('source_ref', 'C:\\private\\source.csv'),
    ('captured_at_utc', '2026-10-26T00:00:00'),
    ('captured_at_utc', '2026-10-26T03:00:00+03:00'),
    ('captured_at_utc', '2026-02-30T00:00:00Z')])
def test_receipt_identity_validation(candidate, cases, field, value):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), **{field: value})
    assert_failure(captured.value, 'A4_INPUT')


@pytest.mark.parametrize('content', [None, 'TRADE_CODE,SUPERTYPE', bytearray(b'not immutable bytes')])
def test_source_input_is_immutable_bytes(candidate, content):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, content)
    assert_failure(captured.value, 'A4_INPUT')


def test_public_default_caps_are_literal(candidate):
    assert candidate.BODY_LIMIT == 16 * 1024**2
    assert candidate.COLUMN_LIMIT == 256
    assert candidate.ROW_LIMIT == 100000
    assert candidate.PARSED_CELL_LIMIT == 1000000
    assert candidate.PAYLOAD_LIMIT == 64 * 1024**2
    assert candidate.CELL_UTF16_LIMIT == 32767


@pytest.mark.parametrize('constant,bound,reason', [
    ('COLUMN_LIMIT', 3, 'A4_BUDGET_COLUMNS'),
    ('ROW_LIMIT', 2, 'A4_BUDGET_ROWS'),
    ('PARSED_CELL_LIMIT', 9, 'A4_BUDGET_PARSED_CELLS')])
def test_small_mechanism_boundary_and_plus_one(candidate, monkeypatch, constant, bound, reason):
    content = encoded(['TRADE_CODE', 'SUPERTYPE', 'EXTRA'],
                      [['ONE', 'Акции', 'a'], ['TWO', 'Облигации', 'b']])
    monkeypatch.setattr(candidate, constant, bound)
    assert prepare(candidate, content)['receipt']['counts']['full_rows'] == 2
    monkeypatch.setattr(candidate, constant, bound - 1)
    with pytest.raises(ValueError) as captured:
        prepare(candidate, content)
    assert_failure(captured.value, reason)


def test_body_boundary_and_plus_one(candidate, monkeypatch, cases):
    content = cases['simple_text'].encode('cp1251')
    monkeypatch.setattr(candidate, 'BODY_LIMIT', len(content))
    assert prepare(candidate, content)['receipt']['counts']['full_rows'] == 1
    monkeypatch.setattr(candidate, 'BODY_LIMIT', len(content) - 1)
    with pytest.raises(ValueError) as captured:
        prepare(candidate, content)
    assert_failure(captured.value, 'A4_BUDGET_BODY_BYTES', 'SOURCE_BODY')


def test_payload_highwater_seek_and_aggregate_boundary(candidate, monkeypatch):
    monkeypatch.setattr(candidate, 'PAYLOAD_LIMIT', 8)
    budget = {'total': 0}
    first = candidate._payload_buffer('moex_full.csv', budget)
    second = candidate._payload_buffer('moex_stocks.csv', budget)
    first.write(b'abcd')
    first.seek(0)
    first.write(b'XY')  # ZIP-style overwrite must not charge the same bytes twice.
    second.write(b'1234')
    assert first.getvalue() == b'XYcd' and second.getvalue() == b'1234'
    assert budget['total'] == 8
    with pytest.raises(ValueError) as captured:
        second.write(b'5')
    assert_failure(captured.value, 'A4_BUDGET_PAYLOAD_BYTES', 'moex_stocks.csv')
    assert budget['total'] == 8 and second.getvalue() == b'1234'
    with pytest.raises(ValueError):
        first.truncate(9)
    assert first.getvalue() == b'XYcd'


def test_real_utf16_cell_boundary_and_plus_one(candidate):
    columns = ['TRADE_CODE', 'SUPERTYPE', 'TEXT']
    # One supplementary Unicode point occupies two UTF-16 units.
    boundary = 'x' * 32765 + '\U0001f642'
    content = encoded(columns, [['SBER', 'Акции', boundary]], 'utf-8')
    result = prepare(candidate, content, content_type='text/csv; charset=utf-8')
    expected = dict(columns=columns, rows=[['SBER', 'Акции', boundary]],
                    stocks=[['SBER', 'Акции', boundary]])
    candidate_oracle(result, expected)
    with pytest.raises(ValueError) as captured:
        prepare(candidate, encoded(columns, [['SBER', 'Акции', boundary + 'x']], 'utf-8'),
                content_type='text/csv; charset=utf-8')
    assert_failure(captured.value, 'A4_XLSX_TEXT_LIMIT')


@pytest.mark.parametrize('value', ['bad\x00xml', 'bad\x0bxml', 'bad\ufffexml'])
def test_unrepresentable_xml_refuses_without_stripping(candidate, value):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, encoded(['TRADE_CODE', 'SUPERTYPE', 'TEXT'], [['SBER', 'Акции', value]], 'utf-8'),
                content_type='text/csv; charset=utf-8')
    assert_failure(captured.value, 'A4_XLSX_VALUE')


@pytest.mark.parametrize('kind', ['disk_full', 'memory', 'render'])
def test_native_render_failure_keeps_primary_and_artifact(candidate, monkeypatch, cases, kind):
    primary = {'disk_full': OSError(errno.ENOSPC, 'synthetic disk full'),
               'memory': MemoryError('synthetic memory'),
               'render': RuntimeError('synthetic exporter')}[kind]

    def fail(frame, artifact, budget):
        raise primary

    monkeypatch.setattr(candidate, '_render_xlsx', fail)
    with pytest.raises(type(primary)) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'))
    assert captured.value is primary
    assert_failure(primary, artifact='moex_full.xlsx')
    if kind == 'disk_full':
        assert primary.errno == errno.ENOSPC


def test_readback_failure_cannot_return_completed_candidate(candidate, monkeypatch, cases):
    primary = ValueError('A4_READBACK')

    def fail(payload, frame, artifact):
        raise primary

    monkeypatch.setattr(candidate, '_validate_xlsx', fail)
    with pytest.raises(ValueError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'))
    assert captured.value is primary
    assert_failure(primary, artifact='moex_full.xlsx')


def mutate_xlsx(payload, mutation):
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(payload), data_only=False)
    try:
        sheet = book['Sheet1']
        if mutation == 'drop_added_column':
            sheet.delete_cols(5)
        elif mutation == 'reorder_rows':
            first = [sheet.cell(2, column).value for column in range(1, 6)]
            second = [sheet.cell(3, column).value for column in range(1, 6)]
            for column, (a, b) in enumerate(zip(first, second), 1):
                sheet.cell(2, column, b)
                sheet.cell(3, column, a)
                if column > 1:
                    sheet.cell(2, column).data_type = sheet.cell(3, column).data_type = 's'
        elif mutation == 'formula':
            sheet['E3'].data_type = 'f'  # Existing literal '=1+1'.
        elif mutation == 'error':
            sheet['D3'].data_type = 'e'  # Existing literal '#N/A'.
        elif mutation == 'extra_sheet':
            book.create_sheet('unexpected')
        else:
            raise AssertionError('Unknown test mutation')
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()
    finally:
        book.close()


@pytest.mark.parametrize('mutation', ['drop_added_column', 'reorder_rows', 'formula', 'error', 'extra_sheet'])
def test_independent_xlsx_oracle_rejects_mutated_candidate(candidate, cases, mutation):
    case = cases['mixed']
    payload = prepare(candidate, encoded(case['columns'], case['rows']))['artifacts']['moex_full.xlsx']
    xlsx_oracle(payload, case['columns'], case['rows'], 1)
    changed = mutate_xlsx(payload, mutation)
    with pytest.raises(AssertionError):
        xlsx_oracle(changed, case['columns'], case['rows'], 1)


@pytest.mark.parametrize('mutation', ['drop_added_column', 'reorder_rows', 'formula', 'error', 'extra_sheet'])
def test_helper_readback_rejects_xlsx_corruption(candidate, cases, monkeypatch, mutation):
    case = cases['mixed']
    original = candidate._render_xlsx

    def corrupt(frame, artifact, budget):
        payload = original(frame, artifact, budget)
        return mutate_xlsx(payload, mutation) if artifact == 'moex_full.xlsx' else payload

    monkeypatch.setattr(candidate, '_render_xlsx', corrupt)
    with pytest.raises(ValueError) as captured:
        prepare(candidate, encoded(case['columns'], case['rows']))
    assert_failure(captured.value, 'A4_READBACK', 'moex_full.xlsx')


@pytest.mark.parametrize('mutation', ['drop_added_column', 'reorder_rows'])
def test_independent_csv_oracle_rejects_mutated_candidate(candidate, cases, mutation):
    case = cases['mixed']
    payload = prepare(candidate, encoded(case['columns'], case['rows']))['artifacts']['moex_full.csv']
    rows = list(csv.reader(io.StringIO(payload.decode('utf-8'), newline='')))
    if mutation == 'drop_added_column':
        rows = [row[:-1] for row in rows]
    else:
        rows[1], rows[2] = rows[2], rows[1]
    stream = io.StringIO(newline='')
    csv.writer(stream).writerows(rows)
    with pytest.raises(AssertionError):
        csv_oracle(stream.getvalue().encode('utf-8'), case['columns'], case['rows'], 1)


def test_only_payload_values_change_with_added_source_field(candidate, cases):
    case = cases['mixed']
    original = encoded(case['columns'], case['rows'])
    changed_rows = [list(row) for row in case['rows']]
    changed_rows[0][-1] = '00124'
    changed = encoded(case['columns'], changed_rows)
    before = prepare(candidate, original)
    after = prepare(candidate, changed)
    assert before['receipt']['source_sha256'] != after['receipt']['source_sha256']
    assert before['artifacts']['moex_full.csv'] != after['artifacts']['moex_full.csv']
    assert before['artifacts']['moex_stocks.csv'] != after['artifacts']['moex_stocks.csv']
    assert before['receipt']['snapshot_id'] == after['receipt']['snapshot_id']


def test_current_targets_are_never_api_arguments_or_written(candidate, cases, tmp_path):
    previous = {name: b'UNVERIFIED_PREVIOUS_' + name.encode('ascii') for name in NAMES}
    for name, content in previous.items():
        (tmp_path / name).write_bytes(content)
    prepare(candidate, cases['simple_text'].encode('cp1251'))
    with pytest.raises(ValueError):
        prepare(candidate, b'not a CSV')
    assert {name: (tmp_path / name).read_bytes() for name in NAMES} == previous
    with pytest.raises(TypeError):
        prepare(candidate, cases['simple_text'].encode('cp1251'), current_path=tmp_path)


@pytest.fixture
def a4_log(monkeypatch):
    # Public transport doubles exercise the real closed envelope/redactor.
    # They do not configure threads, files, queues or a real logging session.
    import run_logging
    context = {'_secrets': ('A4_RAW_CELL_CANARY',), '_checkout': ROOT,
               'config': run_logging.validate_logging_config(), 'run_id': 'a4-fixture',
               'worker_id': None, 'producer_sha': 'fixture-sha', 'package': 'A'}
    state = SimpleNamespace(context=context, events=[], receipts=[], barriers=[],
                            sequence=0, fault=None, only_error=False)

    def faulty(level):
        return not state.only_error or level == 'ERROR'

    def emit(context, level, event, message, fields):
        state.sequence += 1
        if state.fault == 'emit_exception' and faulty(level):
            raise OSError('A4_RAW_CELL_CANARY')
        record = run_logging._event(context, level, event, message, fields)
        state.events.append(record)
        receipt = dict(accepted=True, confirmed=level == 'ERROR', producer=7, sequence=state.sequence)
        if faulty(level):
            if state.fault == 'rejected':
                receipt['accepted'] = False
            elif state.fault == 'producer_drift' and state.sequence > 1:
                receipt['producer'] = 8
            elif state.fault == 'duplicate_sequence' and state.sequence > 1:
                receipt['sequence'] = 1
            elif state.fault == 'boolean_sequence':
                receipt['sequence'] = True
            elif state.fault == 'error_unconfirmed' and level == 'ERROR':
                receipt['confirmed'] = False
        state.receipts.append(receipt)
        return dict(receipt)

    def flush(context):
        state.sequence += 1
        receipt = dict(accepted=True, confirmed=True, producer=7, sequence=state.sequence)
        if state.fault == 'barrier_refused':
            receipt = dict(accepted=False, confirmed=False)
        elif state.fault == 'barrier_boolean':
            receipt['sequence'] = True
        elif state.fault == 'barrier_unconfirmed':
            receipt['confirmed'] = False
        elif state.fault == 'barrier_producer':
            receipt['producer'] = 8
        elif state.fault == 'barrier_stale':
            receipt['sequence'] = 1
        state.barriers.append(receipt)
        return dict(receipt)

    def health(context):
        failed = state.fault == 'unhealthy' and (not state.only_error or
                                               state.events and state.events[-1]['level'] == 'ERROR')
        return {'healthy': not failed, 'state': 'FINALIZED' if state.fault == 'finalized' else 'ACTIVE',
                'code': 'LOGGING_INCOMPLETE' if failed else None}

    monkeypatch.setattr(run_logging, 'emit_event', emit)
    monkeypatch.setattr(run_logging, 'flush_logging', flush)
    monkeypatch.setattr(run_logging, 'check_logging_health', health)
    return state


def test_logging_candidate_requires_same_producer_confirmed_barrier(candidate, cases, a4_log):
    result = prepare(candidate, cases['simple_text'].encode('cp1251'), logging_context=a4_log.context)
    proof = result['receipt']['logging']
    assert proof['status'] == 'CONFIRMED' and proof['producer'] == 7
    assert proof['event_sequences'] == [item['sequence'] for item in a4_log.receipts]
    assert len(proof['event_sequences']) == len(set(proof['event_sequences']))
    assert proof['barrier_sequence'] == a4_log.barriers[-1]['sequence']
    assert proof['barrier_sequence'] > max(proof['event_sequences'])
    assert all(event['fields'] == {} for event in a4_log.events)
    assert all(event['stage'] == 'data_gathering' for event in a4_log.events)
    assert any('ready' in event['event'] for event in a4_log.events)


@pytest.mark.parametrize('fault', ['rejected', 'producer_drift', 'duplicate_sequence', 'boolean_sequence',
                                  'unhealthy', 'finalized', 'barrier_unconfirmed', 'barrier_producer',
                                  'barrier_stale', 'barrier_refused', 'barrier_boolean', 'emit_exception'])
def test_logging_failure_never_returns_candidate(candidate, cases, a4_log, fault):
    a4_log.fault = fault
    with pytest.raises(RuntimeError, match='^A4_LOGGING_INCOMPLETE$'):
        prepare(candidate, cases['simple_text'].encode('cp1251'), logging_context=a4_log.context)


@pytest.mark.parametrize('key', ['console_level', 'file_level'])
def test_logging_requires_info_delivery(candidate, cases, a4_log, monkeypatch, key):
    a4_log.context['config'][key] = 'WARNING'
    rendered = []
    monkeypatch.setattr(candidate, '_render_csv', lambda *args: rendered.append(args))
    with pytest.raises(RuntimeError, match='^A4_LOGGING_INCOMPLETE$'):
        prepare(candidate, cases['simple_text'].encode('cp1251'), logging_context=a4_log.context)
    assert rendered == []


@pytest.mark.parametrize('fault', ['rejected', 'unhealthy', 'emit_exception', 'error_unconfirmed'])
def test_primary_decode_error_survives_secondary_logging_failure(candidate, a4_log, fault):
    a4_log.fault, a4_log.only_error = fault, True
    content = b'TRADE_CODE,SUPERTYPE\r\nSBER,\x98\r\n'
    with pytest.raises(UnicodeDecodeError) as captured:
        prepare(candidate, content, logging_context=a4_log.context)
    assert_failure(captured.value, artifact='SOURCE_BODY')
    assert captured.value.object == content
    assert not any('ready' in event['event'] for event in a4_log.events)


def test_failed_export_logs_fixed_reason_without_raw_cell(candidate, monkeypatch, cases, a4_log, capsys):
    primary = ValueError('A4_RAW_CELL_CANARY C:\\private\\RAW_PATH_CANARY https://user:pass@example.invalid/?token=QUERY_CANARY')

    def fail(frame, artifact, budget):
        raise primary

    monkeypatch.setattr(candidate, '_render_xlsx', fail)
    with pytest.raises(ValueError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), logging_context=a4_log.context)
    assert captured.value is primary
    assert_failure(primary, artifact='moex_full.xlsx')
    output = capsys.readouterr()
    diagnostic = json.dumps(a4_log.events, ensure_ascii=False) + output.out + output.err
    assert all(token not in diagnostic for token in ('A4_RAW_CELL_CANARY', 'RAW_PATH_CANARY', 'QUERY_CANARY', 'user:pass'))
    assert not any('ready' in event['event'] for event in a4_log.events)


def legacy_download(monkeypatch, gathering, content, status):
    calls = []

    class Session:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def get(self, *args, **kwargs):
            calls.append((args, kwargs))
            return SimpleNamespace(content=content, status_code=status,
                                   headers={'Content-Type': 'text/csv; charset=cp1251'})

    monkeypatch.setattr(gathering.requests, 'Session', Session)
    return calls


def test_legacy_red_403_csv_reaches_current_writers(gathering, monkeypatch, writers, fixture_root, cases):
    calls = legacy_download(monkeypatch, gathering, cases['simple_text'].encode('cp1251'), 403)
    gathering.moex_tickerlists(fixture_root[1])
    assert len(calls) == 1
    if os.environ.get('MDS_A4_LEGACY_RED') == '1':
        assert writers == [], 'LEGACY_RED_403: required status refusal before every writer; unchanged caller writes'
    else:
        assert [record[0] for record in writers] == ['xlsx', 'csv', 'csv']


def test_legacy_red_valid_csv_has_only_three_current_writers(gathering, monkeypatch, writers, fixture_root, cases):
    calls = legacy_download(monkeypatch, gathering, cases['simple_text'].encode('cp1251'), 200)
    gathering.moex_tickerlists(fixture_root[1])
    assert len(calls) == 1
    names = [Path(record[2]).name for record in writers]
    if os.environ.get('MDS_A4_LEGACY_RED') == '1':
        assert set(names) == set(NAMES), 'LEGACY_RED_FOUR_FORMATS: required stocks XLSX missing in unchanged caller'
    else:
        assert names == ['moex_full.xlsx', 'moex_full.csv', 'moex_stocks.csv']



def test_payload_budget_refusal_has_no_complete_result(candidate, cases, monkeypatch):
    monkeypatch.setattr(candidate, 'PAYLOAD_LIMIT', 1)
    result = None
    with pytest.raises(ValueError) as captured:
        result = prepare(candidate, cases['simple_text'].encode('cp1251'))
    assert result is None
    assert_failure(captured.value, 'A4_BUDGET_PAYLOAD_BYTES', 'moex_full.csv')


def test_utf16_header_boundary_and_plus_one(candidate):
    boundary = 'x' * 32765 + '\U0001f642'
    rows = [['SBER', 'Акции', 'header value']]
    columns = ['TRADE_CODE', 'SUPERTYPE', boundary]
    result = prepare(candidate, encoded(columns, rows, 'utf-8'), content_type='text/csv; charset=utf-8')
    candidate_oracle(result, dict(columns=columns, rows=rows, stocks=rows))
    with pytest.raises(ValueError) as captured:
        prepare(candidate, encoded(['TRADE_CODE', 'SUPERTYPE', boundary + 'x'], rows, 'utf-8'),
                content_type='text/csv; charset=utf-8')
    assert_failure(captured.value, 'A4_XLSX_TEXT_LIMIT')


def test_failure_log_keeps_bounded_rejected_http_status(candidate, cases, a4_log):
    with pytest.raises(ValueError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), status=403, logging_context=a4_log.context)
    assert_failure(captured.value, 'A4_HTTP_STATUS', 'SOURCE_BODY')
    errors = [item for item in a4_log.events if item['level'] == 'ERROR']
    assert len(errors) == 1 and 'status=403' in errors[0]['message']
    assert errors[0]['error']['code'] == 'A4_HTTP_STATUS'
    assert errors[0]['error']['snapshot_id'] == errors[0]['snapshot_id']
    assert not any('ready' in item['event'] for item in a4_log.events)


def test_forged_export_failure_metadata_never_enters_diagnostics(candidate, cases, a4_log, monkeypatch, capsys):
    primary = OSError(errno.EIO, 'RAW_PATH_CANARY')
    primary.a4_failure = dict(version=1, reason='A4_RAW_PATH_CANARY', artifact='moex_full.xlsx')

    def fail(frame, artifact, budget):
        raise primary

    monkeypatch.setattr(candidate, '_render_xlsx', fail)
    with pytest.raises(OSError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), logging_context=a4_log.context)
    assert captured.value is primary
    assert_failure(primary, 'A4_IO', 'moex_full.xlsx')
    errors = [item for item in a4_log.events if item['level'] == 'ERROR']
    assert len(errors) == 1 and errors[0]['error']['category'] == 'IO'
    assert errors[0]['error']['file'] == 'ticker_lists/moex_full.xlsx'
    output = capsys.readouterr()
    assert 'RAW_PATH_CANARY' not in json.dumps(a4_log.events) + output.out + output.err


@pytest.mark.parametrize('fault', ['rejected', 'unhealthy', 'emit_exception', 'error_unconfirmed'])
def test_primary_render_error_survives_secondary_logging_failure(candidate, cases, a4_log, monkeypatch, fault):
    primary = MemoryError('A4_RAW_CELL_CANARY')
    a4_log.fault, a4_log.only_error = fault, True

    def fail(frame, artifact, budget):
        raise primary

    monkeypatch.setattr(candidate, '_render_xlsx', fail)
    with pytest.raises(MemoryError) as captured:
        prepare(candidate, cases['simple_text'].encode('cp1251'), logging_context=a4_log.context)
    assert captured.value is primary
    assert_failure(primary, 'A4_MEMORY', 'moex_full.xlsx')
    assert not any('ready' in item['event'] for item in a4_log.events)
